from openpyxl import Workbook
from openpyxl.chart import (
    AreaChart,
    Reference,
    Series,
)

class ExcelOper(object):

    def __init__(self):
        self.wb = Workbook()

    def write_excel(self, infos):
        try:
            ws = self.wb.active
            ws.append(['id', 'username', 'funny_num', 'context', 'url'])
            for row_index,row_value in enumerate(infos, 2):
                
                for col_index, col_value in enumerate(row_value, 1):
                    ws.cell(row=row_index, column=col_index, value=col_value)

            self.draw_chart(ws)
            self.wb.save('qiubai.xlsx')
        except Exception as e:
            raise e

    def draw_chart(self, ws):
        chart = AreaChart()
        chart.title = "Joker Chart"
        chart.style = 13
        chart.x_axis.title = 'User'
        chart.y_axis.title = 'Funny Num'

        cats = Reference(ws, min_col=2, min_row=1, max_row=25)
        data = Reference(ws, min_col=3, min_row=1, max_col=3, max_row=25)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        ws.add_chart(chart, "A30")
    def read_excel(self):
        pass